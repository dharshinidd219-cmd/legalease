from flask import Flask, render_template, request, redirect, url_for, flash, send_file
import sqlite3, os, io
from datetime import datetime
from docx import Document

app = Flask(__name__)
app.secret_key = "legalease-demo-key"
DB = "legalease.db"

def db():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con

def init_db():
    con = db()
    con.execute("""CREATE TABLE IF NOT EXISTS documents(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        title TEXT NOT NULL,
        doc_type TEXT NOT NULL,
        content TEXT NOT NULL,
        created_at TEXT NOT NULL
    )""")
    con.commit(); con.close()

TEMPLATES = {
    "Rental Agreement": """RENTAL AGREEMENT

This Rental Agreement is made between {party_a} (Landlord) and {party_b} (Tenant).

Property Address:
{address}

Monthly Rent:
{amount}

Term:
{term}

The Tenant agrees to use the property lawfully and to pay the agreed rent on time. The Landlord agrees to provide peaceful possession subject to the terms of this agreement.

Additional Terms:
{terms}

Date: {date}

Signature – Landlord: ____________________
Signature – Tenant: ______________________
""",
    "Affidavit": """AFFIDAVIT

I, {party_a}, residing at {address}, do hereby solemnly affirm and state that:

{terms}

This affidavit is made for the purpose of providing a written declaration of the above facts.

Date: {date}

Deponent Signature: ______________________
""",
    "NDA": """NON-DISCLOSURE AGREEMENT

This agreement is between {party_a} (Disclosing Party) and {party_b} (Receiving Party).

Purpose:
{terms}

Confidential information shared under this agreement shall be used only for the stated purpose and shall not be disclosed to unauthorized persons, subject to applicable law and the agreed exceptions.

Date: {date}

Disclosing Party: ________________________
Receiving Party: _________________________
""",
    "Employment Agreement": """EMPLOYMENT AGREEMENT

This agreement is between {party_a} (Employer) and {party_b} (Employee).

Role / Position:
{terms}

Salary / Compensation:
{amount}

Term:
{term}

The employee agrees to perform assigned duties responsibly and follow applicable workplace policies. The employer agrees to provide the agreed compensation and working conditions subject to applicable law.

Date: {date}

Employer Signature: ______________________
Employee Signature: ______________________
"""
}

def generate_content(data):
    template = TEMPLATES[data["doc_type"]]
    return template.format(
        party_a=data.get("party_a",""),
        party_b=data.get("party_b",""),
        address=data.get("address",""),
        amount=data.get("amount",""),
        term=data.get("term",""),
        terms=data.get("terms",""),
        date=datetime.now().strftime("%d-%m-%Y")
    )

@app.route("/")
def index():
    con=db(); docs=con.execute("SELECT * FROM documents ORDER BY id DESC").fetchall(); con.close()
    return render_template("index.html", types=TEMPLATES.keys(), docs=docs)

@app.route("/generate", methods=["POST"])
def generate():
    data=request.form.to_dict()
    content=generate_content(data)
    con=db()
    con.execute("INSERT INTO documents(title,doc_type,content,created_at) VALUES(?,?,?,?)",
                (data.get("title") or data["doc_type"], data["doc_type"], content, datetime.now().strftime("%Y-%m-%d %H:%M")))
    con.commit(); con.close()
    return render_template("preview.html", title=data.get("title") or data["doc_type"], content=content)

@app.route("/download/<int:doc_id>")
def download(doc_id):
    con=db(); row=con.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone(); con.close()
    if not row: return "Document not found",404
    doc=Document()
    doc.add_heading(row["title"],0)
    for p in row["content"].split("\n\n"):
        doc.add_paragraph(p)
    path=os.path.join("/tmp", f"legalease_{doc_id}.docx")
    doc.save(path)
    return send_file(path, as_attachment=True, download_name=f"{row['title'].replace(' ','_')}.docx")

@app.route("/history")
def history():
    con=db(); docs=con.execute("SELECT * FROM documents ORDER BY id DESC").fetchall(); con.close()
    return render_template("history.html", docs=docs)

if __name__ == "__main__":
    init_db()
    app.run(debug=True)
